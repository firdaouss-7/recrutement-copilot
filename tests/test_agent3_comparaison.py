# tests/test_agent3_comparaison.py
#
# Etape G du plan Agent 3 : test en isolation + comparaison "avec RAG"
# vs "sans RAG" (ablation study).
#
# Corrections apportees suite a la revue du chat precedent (a faire
# AVANT tout changement d'architecture) :
#
# 1/2. Timing corrige. Le chronometrage initial mesurait le temps total
#    de generer_questions_entretien(), qui englobait — au premier appel
#    "avec RAG" seulement — le chargement paresseux de l'index FAISS et
#    du modele d'embeddings (SentenceTransformer), inexistant cote
#    "sans RAG" puisque le pool est vide et _charger_ressources() n'est
#    jamais appele. L'ecart de temps mesure etait donc en partie un
#    artefact de chargement, pas un cout de retrieval reel. Corrige par :
#      a) un appel de warm-up a _charger_ressources() AVANT toute mesure,
#         chronometre a part et exclu de la comparaison ;
#      b) le decoupage du chronometrage en deux temps distincts pour la
#         version "avec RAG" : temps_retrieval (construire_pool_questions,
#         FAISS + encode) vs temps_generation (l'appel Gemini lui-meme).
#         Cote "sans RAG", seul temps_generation existe (pas de retrieval).
#
# 3. Repete desormais sur plusieurs CV (CVS_TESTES ci-dessous), pas
#    seulement cv2, pour verifier que l'ecart de qualite avec/sans RAG
#    observe n'est pas propre a un seul profil. On garde une diversite
#    de scores/nombre d'ecarts (cv2 : 69/100, 3 ecarts ; cv3 : 50/100,
#    2 ecarts ; cv4 : 78/100, 3 ecarts, deja mentionne en commentaire
#    dans agent3_generation.py comme cas ayant motive nb_questions
#    dynamique). cv1 (0 ecart) reste exclu comme cas extreme deja ecarte
#    dans la version originale.
#
# Sortie : un seul classeur Excel (OUTPUT_PATH), une feuille par CV.
#
# Note (non appliquee ici, cf. echange) : rien dans le code ne garantit
# que chaque point_manquant est reellement couvert dans la sortie finale
# de Gemini, on depend entierement du respect de la consigne du prompt
# (contrairement a l'Agent 2 ou le score est recalcule par du code
# deterministe). Une verification post-generation par similarite
# embedding (point_manquant vs competence_ciblee generees, avec warning
# si aucun match) reste a envisager — pas indispensable pour le PFA,
# mais a mentionner si le jury demande comment la couverture des ecarts
# est garantie.
#
# Usage (depuis la racine du projet) :
#   python -m tests.test_agent3_comparaison

import json
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from agents.agent3_generation import (
    SCHEMA_QUESTIONS_ENTRETIEN,
    _charger_ressources,
    calculer_nb_questions,
    construire_pool_questions,
    construire_prompt,
)
from utils.llm_client import appeler_gemini_json

FIXTURES_DIR = "tests/fixtures/agent2_v2"
CVS_TESTES = ["cv2", "cv3", "cv4"]
OUTPUT_PATH = "tests/fixtures/agent3_comparaison_rag.xlsx"

# --- Palette / styles calques sur tableau_comparaison_agent2.xlsx -------

BLEU_TITRE = "1F4E78"       # fond du titre principal
VERT_SECTION = "D9EAD3"     # fond des en-tetes de section numerotees
BLEU_ENTETE_TABLE = "D9EAF7"  # fond des lignes d'en-tete de tableau

FONT_TITRE = Font(bold=True, size=16, color="FFFFFF")
FONT_SECTION = Font(bold=True, size=12)
FONT_ENTETE_TABLE = Font(bold=True)
FONT_LABEL = Font(bold=True)
FONT_NORMAL = Font()

BORDURE_FINE = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

ALIGN_DONNEE = Alignment(vertical="top", wrap_text=True)
ALIGN_ENTETE_TABLE = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_TITRE = Alignment(horizontal="center", vertical="center")


def charger_json(chemin: str) -> dict:
    with open(chemin, "r", encoding="utf-8") as f:
        return json.load(f)


def executer_test(cle_cv: str, offre: dict, resultat_agent2: dict) -> dict:
    """
    Lance la generation AVEC RAG et SANS RAG pour le CV, et retourne un
    dict regroupant toutes les infos necessaires a l'export Excel.

    Chronometrage (corrige) : suppose que _charger_ressources() a deja
    ete "warm-up" par l'appelant (cf. main()), donc aucun cout de
    chargement FAISS/modele ne doit plus apparaitre ici. On isole :
      - temps_retrieval  : construire_pool_questions() seul (FAISS +
        encode des requetes), UNIQUEMENT cote "avec RAG" ;
      - temps_generation : l'appel Gemini seul (appeler_gemini_json),
        mesure separement pour avec et sans RAG.
    Le "temps total avec RAG" reporte est retrieval + generation ; le
    "temps total sans RAG" est generation seule (pas de retrieval).
    """

    profil = charger_json(f"{FIXTURES_DIR}/{cle_cv}.json")
    nb_questions = calculer_nb_questions(resultat_agent2)
    points_manquants = (resultat_agent2 or {}).get("points_manquants") or []

    print(f"\n=== {cle_cv} ({profil.get('nom')}) ===")
    print(f"score_compatibilite : {resultat_agent2.get('score_compatibilite')}")
    print(f"points_manquants : {points_manquants}")
    print(f"nb_questions calcule : {nb_questions}")

    # --- Retrieval (partage par la version avec RAG uniquement) ---
    debut = time.perf_counter()
    pool_questions = construire_pool_questions(profil, offre, resultat_agent2)
    temps_retrieval = round(time.perf_counter() - debut, 2)
    print(f"Retrieval FAISS : {temps_retrieval}s, {len(pool_questions)} questions dans le pool")

    # --- Avec RAG : generation seule (retrieval deja chronometre) ---
    prompt_avec_rag = construire_prompt(
        profil, offre, pool_questions,
        points_manquants=points_manquants, nb_questions=nb_questions,
    )
    debut = time.perf_counter()
    reponse_avec_rag = appeler_gemini_json(
        prompt_avec_rag,
        model="gemini-3.5-flash-lite",
        thinking_level="medium",
        response_schema=SCHEMA_QUESTIONS_ENTRETIEN,
    )
    temps_generation_avec_rag = round(time.perf_counter() - debut, 2)
    temps_total_avec_rag = round(temps_retrieval + temps_generation_avec_rag, 2)
    print(
        f"Avec RAG : {temps_generation_avec_rag}s generation "
        f"(+{temps_retrieval}s retrieval = {temps_total_avec_rag}s total), "
        f"{len(reponse_avec_rag.get('questions', []))} questions"
    )

    # --- Sans RAG : pas de retrieval, generation seule ---
    prompt_sans_rag = construire_prompt(
        profil, offre, pool_questions=[],
        points_manquants=points_manquants, nb_questions=nb_questions,
    )
    debut = time.perf_counter()
    reponse_sans_rag = appeler_gemini_json(
        prompt_sans_rag,
        model="gemini-3.5-flash-lite",
        thinking_level="medium",
        response_schema=SCHEMA_QUESTIONS_ENTRETIEN,
    )
    temps_generation_sans_rag = round(time.perf_counter() - debut, 2)
    print(
        f"Sans RAG : {temps_generation_sans_rag}s, "
        f"{len(reponse_sans_rag.get('questions', []))} questions"
    )

    return {
        "cv": cle_cv,
        "profil": profil,
        "resultat_agent2": resultat_agent2,
        "nb_questions": nb_questions,
        "temps_retrieval": temps_retrieval,
        "avec_rag": {
            "reponse": reponse_avec_rag,
            "temps_generation": temps_generation_avec_rag,
            "temps_total": temps_total_avec_rag,
        },
        "sans_rag": {
            "reponse": reponse_sans_rag,
            "temps_generation": temps_generation_sans_rag,
            "temps_total": temps_generation_sans_rag,
        },
    }


# --- Construction du classeur Excel -------------------------------------

NB_COLONNES = 5  # N° | Question | Type | Competence ciblee | Justification


def ecrire_titre(ws, ligne: int, texte: str) -> int:
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=NB_COLONNES)
    c = ws.cell(row=ligne, column=1, value=texte)
    c.font = FONT_TITRE
    c.fill = PatternFill("solid", fgColor=BLEU_TITRE)
    c.alignment = ALIGN_TITRE
    ws.row_dimensions[ligne].height = 30
    return ligne + 1


def ecrire_ligne_label_valeur(ws, ligne: int, label: str, valeur) -> int:
    c_label = ws.cell(row=ligne, column=1, value=label)
    c_label.font = FONT_LABEL
    ws.merge_cells(start_row=ligne, start_column=2, end_row=ligne, end_column=NB_COLONNES)
    c_val = ws.cell(row=ligne, column=2, value=valeur)
    c_val.font = FONT_NORMAL
    return ligne + 1


def ecrire_entete_section(ws, ligne: int, texte: str) -> int:
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=NB_COLONNES)
    c = ws.cell(row=ligne, column=1, value=texte)
    c.font = FONT_SECTION
    c.fill = PatternFill("solid", fgColor=VERT_SECTION)
    c.alignment = Alignment(vertical="center")
    return ligne + 1


def ecrire_liste_a_puces(ws, ligne: int, elements: list, symbole: str = "•") -> int:
    if not elements:
        ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=NB_COLONNES)
        c = ws.cell(row=ligne, column=1, value="(aucun ecart identifie)")
        c.font = Font(italic=True, color="808080")
        return ligne + 1

    for item in elements:
        c_symb = ws.cell(row=ligne, column=1, value=symbole)
        c_symb.alignment = ALIGN_DONNEE
        c_symb.border = BORDURE_FINE
        ws.merge_cells(start_row=ligne, start_column=2, end_row=ligne, end_column=NB_COLONNES)
        c_txt = ws.cell(row=ligne, column=2, value=item)
        c_txt.font = FONT_NORMAL
        c_txt.alignment = ALIGN_DONNEE
        c_txt.border = BORDURE_FINE
        ligne += 1
    return ligne


def ecrire_tableau_synthese(ws, ligne: int, resultats: dict) -> int:
    entetes = ["Metrique", "Avec RAG", "Sans RAG", "Ecart", ""]
    for col, texte in enumerate(entetes, start=1):
        c = ws.cell(row=ligne, column=col, value=texte or None)
        c.font = FONT_ENTETE_TABLE
        c.fill = PatternFill("solid", fgColor=BLEU_ENTETE_TABLE)
        c.alignment = ALIGN_ENTETE_TABLE
        c.border = BORDURE_FINE
    ligne += 1

    nb_avec = len(resultats["avec_rag"]["reponse"].get("questions", []))
    nb_sans = len(resultats["sans_rag"]["reponse"].get("questions", []))
    temps_retrieval = resultats["temps_retrieval"]
    temps_gen_avec = resultats["avec_rag"]["temps_generation"]
    temps_gen_sans = resultats["sans_rag"]["temps_generation"]
    temps_total_avec = resultats["avec_rag"]["temps_total"]
    temps_total_sans = resultats["sans_rag"]["temps_total"]

    lignes_donnees = [
        ("Nombre de questions generees", nb_avec, nb_sans, nb_avec - nb_sans),
        ("Nombre de questions demandees", resultats["nb_questions"], resultats["nb_questions"], 0),
        ("Temps retrieval FAISS (s)", temps_retrieval, "—", "n/a"),
        ("Temps generation Gemini (s)", temps_gen_avec, temps_gen_sans, round(temps_gen_avec - temps_gen_sans, 2)),
        ("Temps total (s)", temps_total_avec, temps_total_sans, round(temps_total_avec - temps_total_sans, 2)),
    ]

    for label, val_avec, val_sans, ecart in lignes_donnees:
        valeurs = [label, val_avec, val_sans, ecart, ""]
        for col, val in enumerate(valeurs, start=1):
            c = ws.cell(row=ligne, column=col, value=val if val != "" else None)
            c.font = FONT_LABEL if col == 1 else FONT_NORMAL
            c.alignment = ALIGN_DONNEE if col == 1 else Alignment(horizontal="center", vertical="center")
            c.border = BORDURE_FINE
        ligne += 1

    return ligne


def ecrire_tableau_questions(ws, ligne: int, titre_section: str, questions: list) -> int:
    ligne = ecrire_entete_section(ws, ligne, titre_section)

    entetes = ["N°", "Question", "Type", "Competence ciblee", "Justification"]
    for col, texte in enumerate(entetes, start=1):
        c = ws.cell(row=ligne, column=col, value=texte)
        c.font = FONT_ENTETE_TABLE
        c.fill = PatternFill("solid", fgColor=BLEU_ENTETE_TABLE)
        c.alignment = ALIGN_ENTETE_TABLE
        c.border = BORDURE_FINE
    ligne += 1

    if not questions:
        ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=NB_COLONNES)
        c = ws.cell(row=ligne, column=1, value="(aucune question generee)")
        c.font = Font(italic=True, color="808080")
        return ligne + 1

    for i, q in enumerate(questions, start=1):
        valeurs = [
            i,
            q.get("question", ""),
            q.get("type", ""),
            q.get("competence_ciblee", ""),
            q.get("justification", ""),
        ]
        for col, val in enumerate(valeurs, start=1):
            c = ws.cell(row=ligne, column=col, value=val)
            c.font = FONT_NORMAL
            c.border = BORDURE_FINE
            c.alignment = (
                Alignment(horizontal="center", vertical="top") if col == 1 else ALIGN_DONNEE
            )
        ligne += 1

    return ligne


def construire_feuille(ws, resultats: dict, offre: dict):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Colonnes larges pour eviter le texte tronque (memes proportions
    # que tableau_comparaison_agent2.xlsx : 2 colonnes larges pour le
    # texte libre, colonnes plus etroites pour les champs courts)
    largeurs = {"A": 32, "B": 45, "C": 16, "D": 24, "E": 45}
    for col, largeur in largeurs.items():
        ws.column_dimensions[col].width = largeur

    profil = resultats["profil"]
    resultat_agent2 = resultats["resultat_agent2"]

    ligne = 1
    ligne = ecrire_titre(ws, ligne, "AGENT 3 — COMPARAISON AVEC RAG vs SANS RAG")
    ligne += 1

    # --- Fiche candidat / poste (comme A3-A5 du fichier Agent 2) ---
    ligne = ecrire_ligne_label_valeur(ws, ligne, "Candidat", profil.get("nom", ""))
    ligne = ecrire_ligne_label_valeur(
        ws, ligne, "Poste",
        f"{offre.get('intitule_poste', '')} — {offre.get('entreprise', '')}",
    )
    ligne = ecrire_ligne_label_valeur(
        ws, ligne, "Score Agent 2",
        f"{resultat_agent2.get('score_compatibilite')}/100 ({resultat_agent2.get('niveau_correspondance')})",
    )
    ligne = ecrire_ligne_label_valeur(ws, ligne, "Nb questions demandees", resultats["nb_questions"])
    ligne += 1

    # --- 1. Points manquants (contexte de generation) ---
    ligne = ecrire_entete_section(ws, ligne, "1. POINTS MANQUANTS IDENTIFIES PAR L'AGENT 2")
    ligne = ecrire_liste_a_puces(ws, ligne, resultat_agent2.get("points_manquants") or [], symbole="✗")
    ligne += 1

    # --- 2. Synthese comparative ---
    ligne = ecrire_entete_section(ws, ligne, "2. SYNTHESE COMPARATIVE (VOLUME / TEMPS DE REPONSE)")
    ligne_fige = ligne  # on figera juste apres cette section
    ligne = ecrire_tableau_synthese(ws, ligne, resultats)
    ligne += 1

    # --- 3. Questions avec RAG ---
    ligne = ecrire_tableau_questions(
        ws, ligne, "3. QUESTIONS GENEREES — AVEC RAG",
        resultats["avec_rag"]["reponse"].get("questions", []),
    )
    ligne += 1

    # --- 4. Questions sans RAG ---
    ligne = ecrire_tableau_questions(
        ws, ligne, "4. QUESTIONS GENEREES — SANS RAG",
        resultats["sans_rag"]["reponse"].get("questions", []),
    )

    ws.freeze_panes = f"A{ligne_fige}"


def main():
    offre = charger_json(f"{FIXTURES_DIR}/offre.json")
    resultats_agent2 = charger_json(f"{FIXTURES_DIR}/resultats_agent2.json")

    # --- Warm-up : force le chargement de l'index FAISS et du modele
    # d'embeddings AVANT toute mesure. Sans ca, ce cout (I/O disque +
    # init du SentenceTransformer, potentiellement plusieurs secondes)
    # tombait uniquement sur le premier appel "avec RAG" chronometre et
    # faussait la comparaison. Chronometre a part, hors comparaison.
    debut = time.perf_counter()
    _charger_ressources()
    temps_chargement = round(time.perf_counter() - debut, 2)
    print(f"Chargement FAISS + modele d'embeddings (warm-up, hors comparaison) : {temps_chargement}s")

    wb = Workbook()
    wb.remove(wb.active)  # feuille par defaut vide, remplacee par une feuille par CV

    for cle_cv in CVS_TESTES:
        resultat_agent2_cv = resultats_agent2[cle_cv]
        resultats = executer_test(cle_cv, offre, resultat_agent2_cv)

        ws = wb.create_sheet(title=cle_cv.upper())
        construire_feuille(ws, resultats, offre)

    wb.save(OUTPUT_PATH)
    print(f"\n Comparatif exporte ({len(CVS_TESTES)} CV) : {OUTPUT_PATH}")


if __name__ == "__main__":
    main()